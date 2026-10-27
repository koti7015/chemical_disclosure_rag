"""Build a ~100 row sample Excel file from the full CSV.

The rows are picked on purpose, not at random, so the demo questions have something to find:
every category, discontinued products, the vitamin A family, trade secrets, baby products, one busy company.

Usage:  python scripts/make_sample.py path/to/interviewtestdataset.csv data/sample_100.xlsx
"""
import sys

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

csv_path, out_path = sys.argv[1], sys.argv[2]

DATE_COLS = ["InitialDateReported", "MostRecentDateReported", "DiscontinuedDate",
             "ChemicalCreatedAt", "ChemicalUpdatedAt", "ChemicalDateRemoved"]

df = pd.read_csv(csv_path, low_memory=False)
for c in DATE_COLS:
    df[c] = pd.to_datetime(df[c], format="%m/%d/%Y", errors="coerce")
df = df.drop_duplicates()

parts = [
    df.groupby("PrimaryCategory").sample(4, random_state=1),                               # all 13 categories
    df[df.DiscontinuedDate.notna()].sample(15, random_state=1),                            # discontinued questions
    df[df.ChemicalName.str.contains("retin|vitamin a", case=False, na=False)].sample(10, random_state=1),
    df[df.ChemicalName == "Trade Secret"].sample(4, random_state=1),
    df[df.PrimaryCategory == "Baby Products"].sample(5, random_state=1),
    df[df.ChemicalName.isin(["Talc", "Carbon black"])].sample(6, random_state=1),
    df[df.CompanyName == "New Avon LLC"].sample(8, random_state=1),                        # one busy company
]
sample = (pd.concat(parts).drop_duplicates()
          .sort_values(["CDPHId", "ChemicalId"]).reset_index(drop=True))

with pd.ExcelWriter(out_path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
    sample.to_excel(writer, index=False, sheet_name="records")
    ws = writer.sheets["records"]
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
    for i, col in enumerate(sample.columns, start=1):
        width = min(max(len(str(col)), int(sample[col].astype(str).str.len().quantile(0.9))) + 2, 45)
        ws.column_dimensions[get_column_letter(i)].width = width

print(f"wrote {len(sample)} rows to {out_path}")
